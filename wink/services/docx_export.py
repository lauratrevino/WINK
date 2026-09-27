"""Turns a WINK chat reply (markdown) into a formatted Word document, for
the "Download as Word" button under WINK's answers. Handles the same
markdown the chat renders: headings, tables, bullet and numbered lists,
bold/italic/code, block quotes, horizontal rules, and Mermaid diagrams
(the browser renders each diagram to a PNG and sends it along)."""
import base64
import io
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

NAVY = RGBColor(0x00, 0x28, 0x55)
ORANGE = RGBColor(0xFF, 0x82, 0x00)
MAX_IMAGE_BYTES = 8 * 1024 * 1024

_INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|(?<![\w*])\*[^*\s][^*]*\*(?!\w)|(?<![\w_])_[^_\s][^_]*_(?!\w)|\[[^\]]+\]\([^)]+\))")
_MARKER = re.compile(r"\[\[\s*(?:map|image):[^\]]*\]\]", re.I)


def _add_inline(par, text, bold=False, color=None, size=None):
    text = _MARKER.sub("", text)
    for part in _INLINE.split(text):
        if not part:
            continue
        b, i, code, t = bold, False, False, part
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            t, b = part[2:-2], True
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            t, code = part[1:-1], True
        elif (part.startswith("*") and part.endswith("*")) or (part.startswith("_") and part.endswith("_")):
            if len(part) > 2:
                t, i = part[1:-1], True
        elif part.startswith("[") and "](" in part:
            label, url = part[1:].split("](", 1)
            t = f"{label} ({url[:-1]})"
        run = par.add_run(t)
        run.bold = b or None
        run.italic = i or None
        if code:
            run.font.name = "Consolas"
        if color is not None:
            run.font.color.rgb = color
        if size is not None:
            run.font.size = size


def _shade(cell, hex_fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def _rule(doc):
    p = doc.add_paragraph()
    p_pr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "DDE3F0")):
        bottom.set(qn(k), v)
    bdr.append(bottom)
    p_pr.append(bdr)


def _cells(line):
    return [c.strip() for c in line.strip().strip("|").split("|")]


def build_docx(markdown_text, images=None):
    """Returns (bytes, suggested_filename)."""
    images = list(images or [])
    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(1)
        s.top_margin = s.bottom_margin = Inches(1)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    for lvl, size in ((1, 20), (2, 15), (3, 13), (4, 12)):
        st = doc.styles[f"Heading {lvl}"]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = NAVY

    lines = (markdown_text or "").replace("\r", "").split("\n")
    title = None
    i = 0
    para_buf = []

    def flush():
        if para_buf:
            p = doc.add_paragraph()
            for k, ln in enumerate(para_buf):
                if k:
                    p.add_run().add_break(WD_BREAK.LINE)
                _add_inline(p, ln)
            para_buf.clear()

    is_row = lambda l: bool(re.match(r"^\s*\|.*\|\s*$", l))
    is_div = lambda l: bool(re.match(r"^\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$", l))
    ul = re.compile(r"^(\s*)[-*•]\s+(.+)$")
    ol = re.compile(r"^(\s*)\d+[.)]\s+(.+)$")

    while i < len(lines):
        line = lines[i]
        if not line.strip():
            flush(); i += 1; continue
        if line.strip().startswith("```"):
            flush()
            lang = line.strip()[3:].strip().lower()
            buf = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                buf.append(lines[i]); i += 1
            i += 1
            if lang == "mermaid":
                img = images.pop(0) if images else None
                if img:
                    doc.add_picture(io.BytesIO(img), width=Inches(6.5))
                else:
                    p = doc.add_paragraph()
                    _add_inline(p, "[Diagram: open the chat to view it]", color=RGBColor(0x6B, 0x7A, 0x99))
            else:
                p = doc.add_paragraph()
                r = p.add_run("\n".join(buf))
                r.font.name = "Consolas"
                r.font.size = Pt(9.5)
            continue
        m = re.match(r"^\s*(#{1,4})\s+(.+?)\s*#*\s*$", line)
        if m:
            flush()
            level = len(m.group(1))
            text = m.group(2).replace("**", "")
            if level == 1 and title is None:
                title = text
            h = doc.add_heading(level=level)
            _add_inline(h, text, color=NAVY)
            i += 1; continue
        if re.match(r"^\s*([-*_])(\s*\1){2,}\s*$", line):
            flush(); _rule(doc); i += 1; continue
        if is_row(line) and i + 1 < len(lines) and is_div(lines[i + 1]):
            flush()
            head = _cells(line)
            i += 2
            rows = []
            while i < len(lines) and is_row(lines[i]):
                rows.append(_cells(lines[i])); i += 1
            ncol = len(head)
            table = doc.add_table(rows=1 + len(rows), cols=ncol)
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for c, txt in enumerate(head):
                cell = table.rows[0].cells[c]
                _shade(cell, "002855")
                _add_inline(cell.paragraphs[0], txt, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
            for r, row in enumerate(rows, start=1):
                for c in range(ncol):
                    cell = table.rows[r].cells[c]
                    if r % 2 == 0:
                        _shade(cell, "F4F6FB")
                    _add_inline(cell.paragraphs[0], row[c] if c < len(row) else "")
            doc.add_paragraph()
            continue
        if re.match(r"^\s*>\s?", line):
            flush()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            buf = []
            while i < len(lines) and re.match(r"^\s*>\s?", lines[i]):
                buf.append(re.sub(r"^\s*>\s?", "", lines[i])); i += 1
            _add_inline(p, " ".join(buf), color=RGBColor(0x55, 0x55, 0x55))
            continue
        mu, mo = ul.match(line), ol.match(line)
        if mu or mo:
            flush()
            mm = mu or mo
            indent = len(mm.group(1)) >= 2
            style = "List Bullet" if mu else "List Number"
            if indent:
                style += " 2"
            p = doc.add_paragraph(style=style)
            _add_inline(p, mm.group(2))
            i += 1; continue
        para_buf.append(line)
        i += 1
    flush()

    out = io.BytesIO()
    doc.save(out)
    slug = re.sub(r"[^A-Za-z0-9]+", "-", title or "WINK-document").strip("-")[:60] or "WINK-document"
    return out.getvalue(), f"{slug}.docx"


def decode_images(data_urls):
    """PNG data URLs from the browser -> bytes, skipping anything that
    isn't a PNG or is too large."""
    out, total = [], 0
    for u in (data_urls or [])[:20]:
        if not isinstance(u, str) or not u.startswith("data:image/png;base64,"):
            out.append(None); continue
        try:
            raw = base64.b64decode(u.split(",", 1)[1], validate=True)
        except Exception:
            out.append(None); continue
        total += len(raw)
        if len(raw) > MAX_IMAGE_BYTES or total > 3 * MAX_IMAGE_BYTES or not raw.startswith(b"\x89PNG"):
            out.append(None); continue
        out.append(raw)
    return out
