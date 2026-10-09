"""Turns a WINK chat reply (markdown) into a branded PowerPoint deck, for
the "Download as PowerPoint" button. Convention WINK's prompt asks the
model to follow: '# ' = deck title, each '## ' = one slide, bullets are
slide body, a 'Notes:' line (or block) = speaker notes, markdown tables
become real tables, and Mermaid diagrams arrive as PNGs from the browser."""
import io
import re

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

NAVY = RGBColor(0x00, 0x28, 0x55)
ORANGE = RGBColor(0xFF, 0x82, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x22, 0x2B, 0x3A)
BAND = RGBColor(0xF4, 0xF6, 0xFB)
MAX_SLIDES = 60

_INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*\s][^*]*\*|\[[^\]]+\]\([^)]+\))")
_MARKER = re.compile(r"\[\[\s*(?:map|image):[^\]]*\]\]", re.I)
_NOTES = re.compile(r"^\s*(?:\*\*)?(?:speaker\s+)?notes?(?:\*\*)?\s*:\s*(?:\*\*)?\s*(.*)$", re.I)
_UL = re.compile(r"^(\s*)[-*•]\s+(.+)$")
_OL = re.compile(r"^(\s*)\d+[.)]\s+(.+)$")


def _runs(par, text, size, color=INK, bold=False):
    text = _MARKER.sub("", text)
    for part in _INLINE.split(text):
        if not part:
            continue
        b, i, t = bold, False, part
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            t, b = part[2:-2], True
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            t = part[1:-1]
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            t, i = part[1:-1], True
        elif part.startswith("[") and "](" in part:
            label, url = part[1:].split("](", 1)
            t = f"{label} ({url[:-1]})"
        r = par.add_run()
        r.text = t
        r.font.size = Pt(size)
        r.font.bold = b
        r.font.italic = i
        r.font.color.rgb = color
        r.font.name = "Calibri"


def _rect(slide, x, y, w, h, fill):
    shp = slide.shapes.add_shape(1, x, y, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def _parse(markdown_text):
    """-> (deck_title, [slide dicts])."""
    lines = (markdown_text or "").replace("\r", "").split("\n")
    title, slides, cur = None, [], None
    in_code = False
    code_lang = ""
    i = 0
    is_row = lambda l: bool(re.match(r"^\s*\|.*\|\s*$", l))
    is_div = lambda l: bool(re.match(r"^\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$", l))
    cells = lambda l: [c.strip() for c in l.strip().strip("|").split("|")]
    in_notes = False

    def new(t):
        nonlocal cur, in_notes
        cur = {"title": t, "bullets": [], "notes": [], "table": None, "diagram": False}
        slides.append(cur)
        in_notes = False

    while i < len(lines):
        line = lines[i]
        s = line.strip()
        if s.startswith("```"):
            lang = s[3:].strip().lower()
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                i += 1
            i += 1
            if lang == "mermaid" and cur is not None:
                cur["diagram"] = True
            continue
        m = re.match(r"^\s*(#{1,4})\s+(.+?)\s*#*\s*$", line)
        if m:
            level, text = len(m.group(1)), m.group(2).replace("**", "")
            if level == 1 and title is None:
                title = text
                new(text); cur["cover"] = True
            elif level <= 2:
                new(text)
            else:  # ### becomes a bold sub-heading line inside the slide
                if cur is None:
                    new(title or "Overview")
                cur["bullets"].append((0, "**" + text + "**"))
                in_notes = False
            i += 1; continue
        if not s or re.match(r"^([-*_])(\s*\1){2,}\s*$", s):
            i += 1; continue
        if cur is None:
            new(title or "Overview")
        n = _NOTES.match(line)
        if n:
            in_notes = True
            if n.group(1):
                cur["notes"].append(n.group(1))
            i += 1; continue
        if is_row(line) and i + 1 < len(lines) and is_div(lines[i + 1]):
            head = cells(line); i += 2; rows = []
            while i < len(lines) and is_row(lines[i]):
                rows.append(cells(lines[i])); i += 1
            cur["table"] = [head] + rows
            continue
        mu, mo = _UL.match(line), _OL.match(line)
        if (mu or mo) and not in_notes:
            mm = mu or mo
            cur["bullets"].append((1 if len(mm.group(1)) >= 2 else 0, mm.group(2)))
        elif in_notes:
            cur["notes"].append(re.sub(r"^\s*[-*•]\s+", "", s))
        else:
            cur["bullets"].append((0, re.sub(r"^>\s?", "", s)))
        i += 1
    return title, slides[:MAX_SLIDES]


def build_pptx(markdown_text, images=None):
    """Returns (bytes, suggested_filename)."""
    images = list(images or [])
    title, slides = _parse(markdown_text)
    if not slides:
        slides = [{"title": "WINK", "bullets": [(0, (markdown_text or "").strip()[:600])],
                   "notes": [], "table": None, "diagram": False}]
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    W, H = prs.slide_width, prs.slide_height
    blank = prs.slide_layouts[6]

    for idx, sd in enumerate(slides):
        sl = prs.slides.add_slide(blank)
        if sd.get("cover") or idx == 0 and not sd["bullets"] and not sd["table"]:
            _rect(sl, 0, 0, W, H, NAVY)
            _rect(sl, Inches(0.8), Inches(3.55), Inches(1.6), Inches(0.08), ORANGE)
            tb = sl.shapes.add_textbox(Inches(0.8), Inches(1.9), Inches(11.7), Inches(1.6))
            tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.BOTTOM
            _runs(tf.paragraphs[0], sd["title"], 44, WHITE, True)
            if sd["bullets"]:
                tb2 = sl.shapes.add_textbox(Inches(0.8), Inches(3.9), Inches(11.7), Inches(2.5))
                tf2 = tb2.text_frame; tf2.word_wrap = True
                for k, (_, t) in enumerate(sd["bullets"][:6]):
                    p = tf2.paragraphs[0] if k == 0 else tf2.add_paragraph()
                    _runs(p, t, 20, RGBColor(0xDD, 0xE3, 0xF0))
        else:
            _rect(sl, 0, 0, W, Inches(1.25), NAVY)
            _rect(sl, 0, Inches(1.25), W, Inches(0.06), ORANGE)
            tb = sl.shapes.add_textbox(Inches(0.6), Inches(0.15), Inches(12.1), Inches(1.0))
            tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            _runs(tf.paragraphs[0], sd["title"], 30, WHITE, True)

            body_top, body_h = Inches(1.6), Inches(5.3)
            img = None
            if sd["diagram"]:
                img = images.pop(0) if images else None
            if sd["table"]:
                rows = sd["table"][:14]
                ncol = max(len(r) for r in rows)
                shape = sl.shapes.add_table(len(rows), ncol, Inches(0.6), body_top, Inches(12.1),
                                            Inches(0.45) * len(rows))
                tbl = shape.table
                fs = 14 if len(rows) <= 8 else 11
                for r, row in enumerate(rows):
                    for c in range(ncol):
                        cell = tbl.cell(r, c)
                        cell.fill.solid()
                        cell.fill.fore_color.rgb = NAVY if r == 0 else (BAND if r % 2 == 0 else WHITE)
                        p = cell.text_frame.paragraphs[0]
                        _runs(p, row[c] if c < len(row) else "", fs, WHITE if r == 0 else INK, r == 0)
                body_top = body_top + Inches(0.45) * len(rows) + Inches(0.3)
                body_h = H - body_top - Inches(0.4)
            if img:
                sl.shapes.add_picture(io.BytesIO(img), Inches(0.8), body_top, height=min(body_h, Inches(4.8)))
            elif sd["bullets"] and body_h > Inches(1):
                n = len(sd["bullets"])
                fs = 24 if n <= 4 else 20 if n <= 6 else 17 if n <= 9 else 14
                tb = sl.shapes.add_textbox(Inches(0.7), body_top, Inches(11.9), body_h)
                tf = tb.text_frame; tf.word_wrap = True
                for k, (lvl, t) in enumerate(sd["bullets"][:14]):
                    p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
                    mark = "•  " if lvl == 0 else "–  "
                    p.level = 0
                    p.space_after = Pt(8)
                    if lvl:
                        p.left_indent = Inches(0.4) if hasattr(p, "left_indent") else None
                    bold_head = t.startswith("**") and t.endswith("**") and lvl == 0
                    _runs(p, ("" if bold_head else mark + ("   " if lvl else "")) + t, fs - (3 if lvl else 0),
                          NAVY if bold_head else INK)
            # footer
            ft = sl.shapes.add_textbox(Inches(0.6), H - Inches(0.45), Inches(12.1), Inches(0.35))
            fp = ft.text_frame.paragraphs[0]; fp.alignment = PP_ALIGN.RIGHT
            _runs(fp, f"WINK  |  {idx + 1}", 10, RGBColor(0x6B, 0x7A, 0x99))
        if sd["notes"]:
            sl.notes_slide.notes_text_frame.text = _MARKER.sub("", "\n".join(sd["notes"]))

    out = io.BytesIO()
    prs.save(out)
    slug = re.sub(r"[^A-Za-z0-9]+", "-", title or slides[0]["title"] or "WINK-slides").strip("-")[:60] or "WINK-slides"
    return out.getvalue(), f"{slug}.pptx"
