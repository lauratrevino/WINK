import io
import os
import subprocess
import zipfile

import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from wink import config
from wink.services import documents as docs
from wink.services import vision


@pytest.fixture()
def fake_vision(monkeypatch):
    monkeypatch.setattr(vision, "vision_available", lambda: True)
    monkeypatch.setattr(vision, "_ask", lambda content, max_tokens=2000: "VISION: a diagram of a cell")
    return None


def _png(path, size=(200, 200)):
    Image.new("RGB", size, "red").save(path)


class TestAllowedTypes:
    @pytest.mark.parametrize("ext", ["pptx", "png", "jpg", "jpeg", "gif", "webp", "heic", "mp4", "mov", "webm"])
    def test_extension_allowed(self, ext):
        assert ext in config.ALLOWED_EXT

    def test_no_default_size_limit(self, app):
        assert config.MAX_UPLOAD_BYTES is None
        assert app.config["MAX_CONTENT_LENGTH"] is None


class TestImages:
    def test_photo_is_described_by_vision(self, tmp_path, fake_vision):
        p = tmp_path / "photo.jpg"
        _png(p)
        out = docs.extract_text(str(p), "photo.jpg")
        assert "VISION: a diagram of a cell" in out and "photo.jpg" in out

    def test_falls_back_to_ocr_when_vision_off(self, tmp_path, monkeypatch):
        monkeypatch.setattr(vision, "vision_available", lambda: False)
        p = tmp_path / "photo.png"
        _png(p)
        assert "photo.png" in docs.extract_text(str(p), "photo.png")

    def test_large_photo_is_downscaled_for_api(self):
        big = Image.new("RGB", (6000, 4000), "blue")
        b64 = vision._prepare_image(big)
        import base64
        im = Image.open(io.BytesIO(base64.b64decode(b64)))
        assert max(im.size) <= config.VISION_MAX_EDGE_PX


class TestPowerPoint:
    def _deck(self, tmp_path):
        png = tmp_path / "pic.png"
        _png(png)
        prs = Presentation()
        s = prs.slides.add_slide(prs.slide_layouts[5])
        s.shapes.title.text = "Mitosis"
        s.shapes.add_picture(str(png), Inches(1), Inches(2))
        tbl = s.shapes.add_table(2, 2, Inches(1), Inches(4), Inches(4), Inches(1)).table
        tbl.cell(0, 0).text = "Phase"; tbl.cell(0, 1).text = "Order"
        s.notes_slide.notes_text_frame.text = "Remember PMAT"
        path = tmp_path / "deck.pptx"
        prs.save(path)
        return str(path)

    def test_text_table_notes_and_pictures(self, tmp_path, fake_vision):
        out = docs.extract_text(self._deck(tmp_path), "deck.pptx")
        assert "Mitosis" in out
        assert "Phase | Order" in out
        assert "Remember PMAT" in out
        assert "VISION: a diagram of a cell" in out

    def test_deck_without_vision_still_gives_text(self, tmp_path, monkeypatch):
        monkeypatch.setattr(vision, "vision_available", lambda: False)
        out = docs.extract_text(self._deck(tmp_path), "deck.pptx")
        assert "Mitosis" in out and "VISION" not in out

    def test_large_deck_passes_zip_check(self, tmp_path):
        p = tmp_path / "big.pptx"
        with zipfile.ZipFile(p, "w", zipfile.ZIP_STORED) as zf:
            zf.writestr("ppt/media/a.bin", os.urandom(300 * 1024 * 1024))
        assert docs._zip_bomb_safe(str(p)) is True


@pytest.mark.skipif(subprocess.run(["which", "ffmpeg"], capture_output=True).returncode != 0, reason="ffmpeg missing")
class TestVideo:
    def test_video_frames_are_described(self, tmp_path, fake_vision):
        p = tmp_path / "clip.mp4"
        subprocess.run(["ffmpeg", "-v", "error", "-f", "lavfi", "-i", "testsrc=duration=12:size=320x240:rate=5",
                        "-pix_fmt", "yuv420p", "-y", str(p)], check=True)
        out = docs.extract_text(str(p), "clip.mp4")
        assert "Video file: clip.mp4" in out
        assert "VISION: a diagram of a cell" in out
        assert "audio is not transcribed" in out.lower()
